"""Import the official IKPU catalog Excel export into schema `tasnif`.

    # Dry run (default): parse, validate, print a report. Touches nothing remote.
    uv run --with openpyxl --with certifi scripts/tasnif/import_catalog.py --file ~/Downloads/catalog-excel.xlsx

    # Apply to a Supabase project.
    uv run --with openpyxl --with certifi scripts/tasnif/import_catalog.py --file ... --apply --project-ref hlmcoirjaydrfqcmnuun

Why it works the way it does:

* Transport is the Supabase Management API (`/v1/projects/{ref}/database/query`),
  authenticated with the Supabase CLI token (env `SUPABASE_ACCESS_TOKEN`, else the
  macOS keychain entry the CLI writes). It needs no database password, no service
  key on disk and no PostgREST exposure of `tasnif`, which stays service-role only.

* `kraftabase` also serves Krafta's shops and payments on a small instance, so the
  import never rewrites what hasn't changed. Rows are staged into the UNLOGGED
  `tasnif.import_rows` (no WAL), then merged one catalog group at a time with
  `ON CONFLICT ... DO UPDATE ... WHERE <row differs>`: an unchanged nightly export
  writes almost nothing to `tasnif.codes`.

* Codes missing from the export are marked inactive, never deleted, and only when
  the export is plausibly complete (at least 90% of the currently active count), so
  a truncated download can't switch off the catalog.

* The export is messy in ways that matter: ~186k cells carry stray whitespace or
  newlines, 39 tree nodes have more than one spelling, and digits 12-14 of a code are
  a brand for some rows but a bare numbering slot for ~97k others ("... ---"). The
  parser normalises whitespace, picks each node's most common name, takes brand from
  the brand *name* only, and reports every anomaly instead of hiding it.
"""

from __future__ import annotations

import argparse
import collections
import hashlib
import json
import os
import re
import ssl
import subprocess
import sys
import time
import urllib.error
import urllib.request
from dataclasses import asdict, dataclass
from pathlib import Path

EXPECTED_HEADER = [
    "Группа", "Класс", "Позиция", "Субпозиция", "Бренд", "Атрибут", "ИКПУ", "Название ИКПУ",
    "Штрих код", "Закрепленные (фиксированные)", None, None, "Рекомендованные к использованию", None,
    "ID-льготы",
]
EXPECTED_SUBHEADER = [
    None, None, None, None, None, None, None, None, None,
    "меры измерения", "единицы измерения", "Упаковка", "меры измерения", "единицы измерения", None,
]
NODE_LENGTHS = (3, 5, 8, 11)  # group, class, position, sub-position

WHITESPACE = re.compile(r"\s+")
HIERARCHY_CELL = re.compile(r"^(\d+)-(.*)$", re.S)
BRAND_CELL = re.compile(r"^(\d{14})-(.+)$", re.S)
IKPU = re.compile(r"^\d{17}$")
BARCODE = re.compile(r"^\d{8,14}$")

CODE_FIELDS = [
    "ikpu", "name_ru", "brand_name", "attribute_ru", "barcode", "fixed_measure_ru", "fixed_unit_ru",
    "fixed_package_ru", "recommended_measure_ru", "recommended_unit_ru", "benefit_id",
]
# Columns compared to decide whether an existing code really changed.
COMPARED_FIELDS = CODE_FIELDS[1:]


def clean(value: object) -> str | None:
    """Collapse every whitespace run (tabs, newlines, NBSP) to one space; blank -> None."""
    if value is None:
        return None
    text = WHITESPACE.sub(" ", str(value)).strip()
    return text or None


@dataclass
class CodeRow:
    ikpu: str
    name_ru: str
    brand_name: str | None
    attribute_ru: str | None
    barcode: str | None
    fixed_measure_ru: str | None
    fixed_unit_ru: str | None
    fixed_package_ru: str | None
    recommended_measure_ru: str | None
    recommended_unit_ru: str | None
    benefit_id: str | None


def parse(path: Path) -> tuple[list[CodeRow], dict[str, str], dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)

    header = [clean(c) for c in next(rows)]
    subheader = [clean(c) for c in next(rows)]
    if header[:15] != EXPECTED_HEADER or subheader[:15] != EXPECTED_SUBHEADER:
        raise SystemExit(
            "The export's header changed; refusing to guess the column mapping.\n"
            f"  header:    {header}\n  subheader: {subheader}"
        )

    codes: list[CodeRow] = []
    seen: set[str] = set()
    node_names: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    anomalies = collections.Counter()
    examples: dict[str, list] = collections.defaultdict(list)

    def note(kind: str, example: object) -> None:
        anomalies[kind] += 1
        if len(examples[kind]) < 5:
            examples[kind].append(example)

    for raw in rows:
        cells = list(raw[:15]) + [None] * (15 - len(raw[:15]))
        ikpu = clean(cells[6])
        if ikpu is None and all(clean(c) is None for c in cells):
            continue
        if ikpu is None or not IKPU.match(ikpu):
            note("invalid_ikpu", cells[6])
            continue
        if ikpu in seen:
            note("duplicate_ikpu", ikpu)
            continue
        seen.add(ikpu)

        for index, length in enumerate(NODE_LENGTHS):
            match = HIERARCHY_CELL.match(str(cells[index] or ""))
            name = clean(match.group(2)) if match else None
            if not match or match.group(1) != ikpu[:length] or not name:
                note(f"bad_hierarchy_cell_len{length}", (ikpu, cells[index]))
                continue
            node_names[ikpu[:length]][name] += 1

        name_ru = clean(cells[7])
        if name_ru is None:
            note("missing_name", ikpu)
            continue

        brand_cell = str(cells[4] or "").strip()
        brand_match = BRAND_CELL.match(brand_cell)
        brand_name = clean(brand_match.group(2)) if brand_match else None
        if brand_match and brand_match.group(1) != ikpu[:14]:
            note("brand_code_not_prefix_of_ikpu", (ikpu, brand_cell))
        elif not brand_match and re.fullmatch(r"\d{14}-", brand_cell):
            note("brand_code_without_name", (ikpu, brand_cell))
        elif not brand_match and not brand_cell.endswith("---"):
            note("unreadable_brand_cell", (ikpu, brand_cell))

        attribute = clean(cells[5])
        barcode = clean(cells[8])
        if barcode is not None and not BARCODE.match(barcode):
            note("invalid_barcode", (ikpu, barcode))
            barcode = None
        benefit = clean(cells[14])

        codes.append(CodeRow(
            ikpu=ikpu,
            name_ru=name_ru,
            brand_name=brand_name,
            attribute_ru=None if attribute == "---" else attribute,
            barcode=barcode,
            fixed_measure_ru=clean(cells[9]),
            fixed_unit_ru=clean(cells[10]),
            fixed_package_ru=clean(cells[11]),
            recommended_measure_ru=clean(cells[12]),
            recommended_unit_ru=clean(cells[13]),
            benefit_id=None if benefit in (None, "0") else benefit,
        ))

    nodes: dict[str, str] = {}
    for code, names in node_names.items():
        nodes[code] = names.most_common(1)[0][0]
        if len(names) > 1:
            note("node_with_several_names", {code: dict(names)})

    for row in codes:
        for length in NODE_LENGTHS:
            if row.ikpu[:length] not in nodes:
                raise SystemExit(f"Code {row.ikpu} has no named ancestor {row.ikpu[:length]}; the export is incomplete.")

    return codes, nodes, {"anomalies": dict(anomalies), "examples": dict(examples)}


def report(path: Path, digest: str, codes: list[CodeRow], nodes: dict[str, str], issues: dict) -> dict:
    kinds = collections.Counter(
        "catering" if c.ikpu.startswith("10202") else "service" if c.ikpu.startswith("1") else "goods"
        for c in codes
    )
    summary = {
        "file": str(path),
        "sha256": digest,
        "codes": len(codes),
        "nodes": {name: sum(1 for k in nodes if len(k) == length)
                  for name, length in zip(("groups", "classes", "positions", "subpositions"), NODE_LENGTHS)},
        "kinds": dict(kinds),
        "branded": sum(1 for c in codes if c.brand_name),
        "with_barcode": sum(1 for c in codes if c.barcode),
        "with_benefit": sum(1 for c in codes if c.benefit_id),
        "anomalies": issues["anomalies"],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if issues["examples"]:
        print("examples:")
        for kind, sample in issues["examples"].items():
            print(f"  {kind}: {json.dumps(sample, ensure_ascii=False)[:400]}")
    return summary


class ManagementApi:
    def __init__(self, project_ref: str) -> None:
        self.url = f"https://api.supabase.com/v1/projects/{project_ref}/database/query"
        self.token = os.environ.get("SUPABASE_ACCESS_TOKEN") or self._keychain_token()
        try:
            import certifi
            self.context = ssl.create_default_context(cafile=certifi.where())
        except ImportError:
            self.context = ssl.create_default_context()

    @staticmethod
    def _keychain_token() -> str:
        # The Supabase CLI stores its token as "go-keyring-base64:<base64>".
        stored = subprocess.run(
            ["security", "find-generic-password", "-s", "Supabase CLI", "-w"],
            capture_output=True, text=True,
        ).stdout.strip()
        if not stored:
            raise SystemExit("No SUPABASE_ACCESS_TOKEN and no Supabase CLI token in the keychain. Run `supabase login`.")
        import base64
        return base64.b64decode(stored.removeprefix("go-keyring-base64:")).decode()

    def query(self, sql: str, attempts: int = 6) -> list[dict]:
        body = json.dumps({"query": sql}).encode()
        for attempt in range(1, attempts + 1):
            request = urllib.request.Request(self.url, data=body, method="POST", headers={
                "Authorization": f"Bearer {self.token}",
                "Content-Type": "application/json",
                "User-Agent": "tasnif-import (tasnif.krafta.uz)",
            })
            try:
                with urllib.request.urlopen(request, timeout=300, context=self.context) as response:
                    return json.loads(response.read() or b"[]")
            except urllib.error.HTTPError as error:
                detail = error.read().decode("utf-8", "replace")[:1000]
                retryable = error.code == 429 or error.code >= 500
                if not retryable or attempt == attempts:
                    raise RuntimeError(f"HTTP {error.code}: {detail}") from None
            except (urllib.error.URLError, TimeoutError) as error:
                if attempt == attempts:
                    raise RuntimeError(f"network: {error}") from None
            time.sleep(min(60, 2 ** attempt))
        raise AssertionError("unreachable")


def sql_json(payload: object) -> str:
    """A jsonb SQL literal. standard_conforming_strings is on, so only quotes need doubling."""
    return "'" + json.dumps(payload, ensure_ascii=False).replace("'", "''") + "'::jsonb"


def batched(items: list, size: int):
    for start in range(0, len(items), size):
        yield items[start:start + size]


def apply(api: ManagementApi, path: Path, digest: str, codes: list[CodeRow], nodes: dict[str, str],
          issues: dict, batch_size: int, pause: float, only_groups: set[str] | None) -> None:
    if only_groups:
        codes = [c for c in codes if c.ikpu[:3] in only_groups]
        nodes = {k: v for k, v in nodes.items() if k[:3] in only_groups}

    active_before = api.query("select count(*)::int as n from tasnif.codes where status = 'active'")[0]["n"]
    notes = {"anomalies": issues["anomalies"], "only_groups": sorted(only_groups) if only_groups else None}
    run_id = api.query(
        "insert into tasnif.sync_runs (source, source_file, source_sha256, rows_seen, notes) values "
        f"('excel', {sql_text(path.name)}, {sql_text(digest)}, {len(codes)}, {sql_json(notes)}) returning id"
    )[0]["id"]
    print(f"sync run {run_id}: {len(codes)} codes, {len(nodes)} nodes, {active_before} active before")

    try:
        # Parents first: the tree's foreign keys point upwards.
        node_added = node_changed = 0
        ordered = sorted(nodes.items(), key=lambda item: (len(item[0]), item[0]))
        for chunk in batched([{"code": k, "name_ru": v} for k, v in ordered], 2000):
            result = api.query(f"""
                with incoming as (
                  select * from jsonb_to_recordset({sql_json(chunk)}) as t(code text, name_ru text)
                ), written as (
                  insert into tasnif.nodes as n (code, name_ru)
                  select code, name_ru from incoming
                  on conflict (code) do update set name_ru = excluded.name_ru, updated_at = now()
                  where n.name_ru is distinct from excluded.name_ru
                  returning (xmax = 0) as inserted
                )
                select count(*) filter (where inserted)::int as added,
                       count(*) filter (where not inserted)::int as changed
                from written""")
            node_added += result[0]["added"]
            node_changed += result[0]["changed"]
            time.sleep(pause)
        print(f"nodes: {node_added} added, {node_changed} renamed")

        columns = ", ".join(f"{field} text" for field in CODE_FIELDS)
        for number, chunk in enumerate(batched([asdict(c) for c in codes], batch_size), start=1):
            api.query(f"""
                insert into tasnif.import_rows (run_id, {", ".join(CODE_FIELDS)})
                select {run_id}, {", ".join(CODE_FIELDS)}
                from jsonb_to_recordset({sql_json(chunk)}) as t({columns})
                on conflict (run_id, ikpu) do nothing""")
            if number % 20 == 0:
                print(f"  staged {min(number * batch_size, len(codes))}/{len(codes)}")
            time.sleep(pause)

        added = changed = 0
        differs = " or ".join(f"c.{field} is distinct from excluded.{field}" for field in COMPARED_FIELDS)
        updates = ", ".join(f"{field} = excluded.{field}" for field in COMPARED_FIELDS)
        groups = sorted({c.ikpu[:3] for c in codes})
        for group in groups:
            result = api.query(f"""
                with written as (
                  insert into tasnif.codes as c ({", ".join(CODE_FIELDS)})
                  select {", ".join(CODE_FIELDS)} from tasnif.import_rows
                  where run_id = {run_id} and left(ikpu, 3) = '{group}'
                  on conflict (ikpu) do update set {updates},
                    status = 'active', inactive_since = null, updated_at = now()
                  where {differs} or c.status <> 'active'
                  returning (xmax = 0) as inserted
                )
                select count(*) filter (where inserted)::int as added,
                       count(*) filter (where not inserted)::int as changed
                from written""")
            added += result[0]["added"]
            changed += result[0]["changed"]
            time.sleep(pause)
        print(f"codes: {added} added, {changed} changed")

        deactivated = 0
        if only_groups:
            print("partial run: skipping deactivation")
        elif len(codes) < 0.9 * active_before:
            notes["deactivation_skipped"] = f"export has {len(codes)} codes, fewer than 90% of {active_before} active"
            print(notes["deactivation_skipped"])
        else:
            deactivated = api.query(f"""
                with gone as (
                  update tasnif.codes c set status = 'inactive', inactive_since = now(), updated_at = now()
                  where c.status = 'active'
                    and not exists (select 1 from tasnif.import_rows r where r.run_id = {run_id} and r.ikpu = c.ikpu)
                  returning 1
                ) select count(*)::int as n from gone""")[0]["n"]
            print(f"codes: {deactivated} deactivated")

        api.query(f"delete from tasnif.import_rows where run_id = {run_id}")
        api.query(f"""
            update tasnif.sync_runs set finished_at = now(), rows_added = {added}, rows_changed = {changed},
              rows_deactivated = {deactivated},
              notes = notes || {sql_json({"nodes_added": node_added, "nodes_renamed": node_changed, **notes})}
            where id = {run_id}""")
    except Exception as error:
        api.query(f"update tasnif.sync_runs set finished_at = now(), error = {sql_text(str(error)[:2000])} where id = {run_id}")
        raise


def sql_text(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--apply", action="store_true", help="write to the database (default: dry run)")
    parser.add_argument("--project-ref", default="hlmcoirjaydrfqcmnuun")
    parser.add_argument("--batch-size", type=int, default=2000)
    parser.add_argument("--pause", type=float, default=0.3, help="seconds between database requests")
    parser.add_argument("--only-groups", help="comma-separated 3-digit groups, for a partial smoke run")
    args = parser.parse_args()

    path = args.file.expanduser()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    started = time.time()
    codes, nodes, issues = parse(path)
    report(path, digest, codes, nodes, issues)
    print(f"parsed in {time.time() - started:.1f}s")

    if not args.apply:
        print("dry run: nothing written (pass --apply to import)")
        return
    only_groups = {g.strip() for g in args.only_groups.split(",")} if args.only_groups else None
    apply(ManagementApi(args.project_ref), path, digest, codes, nodes, issues, args.batch_size, args.pause, only_groups)
    print(f"done in {time.time() - started:.1f}s")


if __name__ == "__main__":
    main()
