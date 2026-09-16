"""Import official Uzbek names (Latin and Cyrillic) for every code and tree node.

    pnpm --filter tasnif catalog:names --latn ~/Downloads/category_0_lat.xlsx --cyrl ~/Downloads/category_0_uz.xlsx
    pnpm --filter tasnif catalog:names --latn ... --cyrl ... --apply

The files are the same public homepage export as the Russian catalog, requested
with `lang=uz_latn` (category_0_lat.xlsx) and `lang=uz_cyrl` (category_0_uz.xlsx).

Why it works the way it does:

* It only writes names. Russian names, units, barcodes and the code list itself
  come from the Russian export (import_catalog.py), which runs first; a code that
  isn't in tasnif.codes yet is counted and skipped, never inserted from here.
* Both scripts are read before anything is written and applied in one pass, so
  each row is rewritten once per run, not once per language. After the first run
  an unchanged export writes nothing: every update is guarded by
  `is distinct from`.
* The headers are checked column by column. If the tax committee renames or
  reorders them, the script stops instead of writing Cyrillic into the Latin
  column.
"""

from __future__ import annotations

import argparse
import collections
import hashlib
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from import_catalog import (  # noqa: E402
    HIERARCHY_CELL, IKPU, NODE_LENGTHS, ManagementApi, batched, clean, sql_json, sql_text,
)

# Only the columns this script reads are checked.
EXPECTED = {
    "uz_latn": {0: "GURUH NOMI", 1: "SINF NOMI", 2: "POZITSIYA NOMI", 3: "SUBPOZITSIYA NOMI", 6: "MXIK KODI", 7: "MXIK NOMI"},
    "uz_cyrl": {0: "ГУРУҲ НОМИ", 1: "СИНФ НОМИ", 2: "ПОЗИЦИЯ НОМИ", 3: "СУБПОЗИЦИЯ НОМИ", 6: "МХИК КОДИ", 7: "МХИК НОМИ"},
}


def parse(path: Path, lang: str) -> tuple[dict[str, str], dict[str, str], collections.Counter]:
    from openpyxl import load_workbook

    rows = load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)
    header = [clean(c) for c in next(rows)]
    next(rows)  # sub-header (unit block labels), unused here
    for index, label in EXPECTED[lang].items():
        if index >= len(header) or header[index] != label:
            raise SystemExit(f"{path.name}: column {index} is {header[index] if index < len(header) else None!r}, "
                             f"expected {label!r} for {lang}; refusing to guess.")

    codes: dict[str, str] = {}
    node_names: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    anomalies = collections.Counter()
    for raw in rows:
        ikpu, name = clean(raw[6]), clean(raw[7])
        if ikpu is None or not IKPU.match(ikpu):
            if any(clean(c) for c in raw):
                anomalies["invalid_ikpu"] += 1
            continue
        if name is None:
            anomalies["missing_name"] += 1
        elif codes.setdefault(ikpu, name) != name:
            anomalies["duplicate_ikpu_other_name"] += 1
        for index, length in enumerate(NODE_LENGTHS):
            match = HIERARCHY_CELL.match(str(raw[index] or ""))
            node_name = clean(match.group(2)) if match else None
            if match and match.group(1) == ikpu[:length] and node_name:
                node_names[ikpu[:length]][node_name] += 1
            else:
                anomalies[f"bad_hierarchy_cell_len{length}"] += 1
    nodes = {code: names.most_common(1)[0][0] for code, names in node_names.items()}
    anomalies["nodes_with_several_names"] = sum(1 for names in node_names.values() if len(names) > 1)
    return codes, nodes, anomalies


def merge(latn: dict[str, str], cyrl: dict[str, str]) -> list[dict]:
    return [{"key": key, "latn": latn.get(key), "cyrl": cyrl.get(key)} for key in sorted(latn.keys() | cyrl.keys())]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--latn", required=True, type=Path, help="category export with lang=uz_latn")
    parser.add_argument("--cyrl", required=True, type=Path, help="category export with lang=uz_cyrl")
    parser.add_argument("--apply", action="store_true", help="write to the database (default: dry run)")
    parser.add_argument("--project-ref", default="hlmcoirjaydrfqcmnuun")
    parser.add_argument("--pause", type=float, default=0.3, help="seconds between database requests")
    args = parser.parse_args()

    started = time.time()
    latn_codes, latn_nodes, latn_issues = parse(args.latn.expanduser(), "uz_latn")
    cyrl_codes, cyrl_nodes, cyrl_issues = parse(args.cyrl.expanduser(), "uz_cyrl")
    print(f"uz_latn: {len(latn_codes)} codes, {len(latn_nodes)} nodes, anomalies {dict(latn_issues)}")
    print(f"uz_cyrl: {len(cyrl_codes)} codes, {len(cyrl_nodes)} nodes, anomalies {dict(cyrl_issues)}")
    if latn_codes.keys() != cyrl_codes.keys():
        print(f"warning: the two files disagree on {len(latn_codes.keys() ^ cyrl_codes.keys())} codes")
    print(f"parsed in {time.time() - started:.1f}s")
    if not args.apply:
        print("dry run: nothing written (pass --apply to import)")
        return

    api = ManagementApi(args.project_ref)
    digest = hashlib.sha256(args.latn.expanduser().read_bytes() + args.cyrl.expanduser().read_bytes()).hexdigest()
    run_id = api.query(
        "insert into tasnif.sync_runs (source, source_file, source_sha256, rows_seen) values "
        f"('excel_uz', {sql_text(args.latn.name + ' + ' + args.cyrl.name)}, '{digest}', {len(latn_codes)}) returning id"
    )[0]["id"]

    totals = collections.Counter()
    try:
        for table, key, rows in (("nodes", "code", merge(latn_nodes, cyrl_nodes)),
                                 ("codes", "ikpu", merge(latn_codes, cyrl_codes))):
            for number, chunk in enumerate(batched(rows, 5000, max_bytes=500_000), start=1):
                result = api.query(f"""
                    with incoming as (
                      select * from jsonb_to_recordset({sql_json(chunk)}) as i(key text, latn text, cyrl text)
                    ), written as (
                      update tasnif.{table} t
                      set name_uz_latn = coalesce(i.latn, t.name_uz_latn),
                          name_uz_cyrl = coalesce(i.cyrl, t.name_uz_cyrl),
                          updated_at = now()
                      from incoming i
                      where t.{key} = i.key
                        and (t.name_uz_latn is distinct from coalesce(i.latn, t.name_uz_latn)
                             or t.name_uz_cyrl is distinct from coalesce(i.cyrl, t.name_uz_cyrl))
                      returning 1
                    )
                    select (select count(*) from written)::int as updated,
                           (select count(*) from incoming i
                            where not exists (select 1 from tasnif.{table} t where t.{key} = i.key))::int as unknown""")
                totals[f"{table}_updated"] += result[0]["updated"]
                totals[f"{table}_unknown"] += result[0]["unknown"]
                if number % 20 == 0:
                    print(f"  {table}: {dict(totals)}", flush=True)
                time.sleep(args.pause)
            print(f"{table}: {totals[table + '_updated']} updated, {totals[table + '_unknown']} not in tasnif.{table}", flush=True)
        api.query(f"""update tasnif.sync_runs set finished_at = now(), rows_changed = {totals['codes_updated']},
            notes = {sql_json({**totals, 'latn_anomalies': latn_issues, 'cyrl_anomalies': cyrl_issues})}
            where id = {run_id}""")
    except Exception as error:
        api.query(f"update tasnif.sync_runs set finished_at = now(), error = {sql_text(str(error)[:2000])} where id = {run_id}")
        raise
    print(f"done in {time.time() - started:.1f}s")


if __name__ == "__main__":
    main()
