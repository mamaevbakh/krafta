"""Import every code's package codes from the catalog page's units export.

    pnpm --filter tasnif catalog:packages --file ~/Downloads/package_ru.xlsx
    pnpm --filter tasnif catalog:packages --file ~/Downloads/package_ru.xlsx --apply

A package code is the second number a fiscal receipt needs next to the IKPU:
10202001010000002 (cafe coffee drinks) -> 1522997 "шт. (кружка/стакан)".

Where the file comes from: tasnif.soliq.uz/catalog -> «Выгрузить в» -> units ->
«Скачать». That dialog asks for a captcha, so a person downloads it; nothing
here fetches it automatically, and nothing should go around the captcha.

Why it works the way it does:

* The export has two sheets that mean different things, and both are kept:
  «Закрепленные единицы и упаковки» (units the tax committee fixed, origin
  'fixed') and «Сформированные упаковки другими пользователями» (packages other
  businesses created, origin 'user'). Sheets are recognised by their title row,
  not their position, and an unknown sheet stops the import.
* ~576k rows are staged into the UNLOGGED tasnif.import_packages, then merged one
  catalog group at a time, writing only rows that are new or differ. The shared
  instance also serves Krafta's shops and payments.
* A package that disappeared from a complete export is deleted. "Complete" means
  the file has at least 90% as many rows as are stored now, so a truncated
  download can't wipe the table.
* Codes in the file that aren't in tasnif.codes yet (new codes the catalog export
  hasn't caught up with) are counted and skipped; the next catalog import brings
  them in and the next packages import attaches their packages.
"""

from __future__ import annotations

import argparse
import collections
import hashlib
import re
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from import_catalog import IKPU, batched, clean, database, sql_json, sql_text  # noqa: E402

SHEETS = {
    "Закрепленные единицы и упаковки": "fixed",
    "Сформированные упаковки другими пользователями": "user",
}
HEADER = ["ИКПУ", "Название ИКПУ", "Название единицы измерения", "Код единицы измерения"]
PACKAGE_CODE = re.compile(r"^\d{1,18}$")


def parse(path: Path) -> tuple[list[dict], collections.Counter]:
    from openpyxl import load_workbook

    packages: dict[tuple[str, int], dict] = {}
    counts = collections.Counter()
    for sheet in load_workbook(path, read_only=True, data_only=True).worksheets:
        rows = sheet.iter_rows(values_only=True)
        title = clean((next(rows) or [None])[0])
        if title not in SHEETS:
            raise SystemExit(f"{path.name}: sheet {sheet.title!r} starts with {title!r}; expected one of {list(SHEETS)}.")
        origin = SHEETS[title]
        header = [clean(c) for c in next(rows)][:4]
        if header != HEADER:
            raise SystemExit(f"{path.name}: sheet {sheet.title!r} header is {header}, expected {HEADER}.")
        for raw in rows:
            # Columns: ИКПУ, Название ИКПУ (not needed), package name, package code.
            cells = (list(raw) + [None] * 4)[:4]
            ikpu, name, code = clean(cells[0]), clean(cells[2]), clean(cells[3])
            if ikpu is None and name is None and code is None:
                continue
            if ikpu is None or not IKPU.match(ikpu) or code is None or not PACKAGE_CODE.match(code) or name is None:
                counts[f"{origin}_unreadable_row"] += 1
                continue
            key = (ikpu, int(code))
            if key in packages:
                counts["duplicate_package"] += 1
                if packages[key]["origin"] != origin:
                    counts["duplicate_package_in_both_sheets"] += 1
                continue
            packages[key] = {"ikpu": ikpu, "package_code": int(code), "name_ru": name, "origin": origin}
            counts[f"{origin}_rows"] += 1
    return list(packages.values()), counts


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--apply", action="store_true", help="write to the database (default: dry run)")
    parser.add_argument("--project-ref", default="hlmcoirjaydrfqcmnuun")
    parser.add_argument("--pause", type=float, default=0.3, help="seconds between database requests")
    args = parser.parse_args()

    path = args.file.expanduser()
    started = time.time()
    packages, counts = parse(path)
    print(f"{len(packages)} packages for {len({p['ikpu'] for p in packages})} codes; {dict(counts)}")
    print(f"parsed in {time.time() - started:.1f}s")
    if not args.apply:
        print("dry run: nothing written (pass --apply to import)")
        return

    api = database(args.project_ref)
    stored_before = api.query("select count(*)::int as n from tasnif.packages")[0]["n"]
    api.query("delete from tasnif.import_packages")  # one import at a time; leftovers belong to a dead run
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    run_id = api.query(
        "insert into tasnif.sync_runs (source, source_file, source_sha256, rows_seen, notes) values "
        f"('packages_excel', {sql_text(path.name)}, '{digest}', {len(packages)}, {sql_json(counts)}) returning id"
    )[0]["id"]
    print(f"sync run {run_id}: {stored_before} packages stored before")

    totals = collections.Counter()
    try:
        staged = 0
        for number, chunk in enumerate(batched(packages, 10000, max_bytes=600_000), start=1):
            api.query(f"""
                insert into tasnif.import_packages (run_id, ikpu, package_code, name_ru, origin)
                select {run_id}, ikpu, package_code, name_ru, origin
                from jsonb_to_recordset({sql_json(chunk)}) as t(ikpu text, package_code bigint, name_ru text, origin text)
                on conflict do nothing""")
            staged += len(chunk)
            if number % 10 == 0:
                print(f"  staged {staged}/{len(packages)}", flush=True)
            time.sleep(args.pause)

        complete = len(packages) >= 0.9 * stored_before
        if not complete:
            print(f"file has {len(packages)} rows, under 90% of {stored_before} stored: not deleting anything")
        for group in sorted({p["ikpu"][:3] for p in packages}):
            upper = f"{int(group) + 1:03d}"
            result = api.query(f"""
                with source as (
                  select s.ikpu, s.package_code, s.name_ru, s.origin
                  from tasnif.import_packages s
                  where s.run_id = {run_id} and s.ikpu >= '{group}' and s.ikpu < '{upper}'
                ), known as (
                  select s.* from source s where exists (select 1 from tasnif.codes c where c.ikpu = s.ikpu)
                ), written as (
                  insert into tasnif.packages as p (ikpu, package_code, name_ru, origin, fetched_at)
                  select ikpu, package_code, name_ru, origin, now() from known
                  on conflict (ikpu, package_code) do update
                    set name_ru = excluded.name_ru, origin = excluded.origin, fetched_at = now()
                  where p.name_ru is distinct from excluded.name_ru or p.origin is distinct from excluded.origin
                  returning (xmax = 0) as inserted
                ), removed as (
                  delete from tasnif.packages p
                  where {str(complete).lower()}
                    and p.ikpu >= '{group}' and p.ikpu < '{upper}'
                    and not exists (select 1 from source s where s.ikpu = p.ikpu and s.package_code = p.package_code)
                  returning 1
                )
                select (select count(*) from written where inserted)::int as added,
                       (select count(*) from written where not inserted)::int as changed,
                       (select count(*) from removed)::int as removed,
                       (select count(*) from source)::int - (select count(*) from known)::int as unknown_code""")
            for key, value in result[0].items():
                totals[key] += value
            time.sleep(args.pause)
        print(f"packages: {dict(totals)}")

        api.query(f"delete from tasnif.import_packages where run_id = {run_id}")
        api.query(f"""update tasnif.sync_runs set finished_at = now(), rows_added = {totals['added']},
            rows_changed = {totals['changed']}, rows_deactivated = {totals['removed']},
            notes = notes || {sql_json({'unknown_code': totals['unknown_code'], 'deleted_missing': complete})}
            where id = {run_id}""")
    except Exception as error:
        api.query(f"delete from tasnif.import_packages where run_id = {run_id}")
        api.query(f"update tasnif.sync_runs set finished_at = now(), error = {sql_text(str(error)[:2000])} where id = {run_id}")
        raise
    print(f"done in {time.time() - started:.1f}s")


if __name__ == "__main__":
    main()
