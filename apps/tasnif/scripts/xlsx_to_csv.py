"""Convert an official catalog export (.xlsx) to one CSV per sheet, for reading and analysis.

    uv run --with openpyxl scripts/xlsx_to_csv.py .data/exports/category_0_ru.xlsx .data/csv/ru

Writes <out_dir>/<sheet name>.csv exactly as the cells are, header rows included (the catalog
exports have two). Nothing is cleaned here: the importers do that, and reviewers should see what
the tax committee actually published. Streams the workbook, so the 30-40 MB exports convert in
under half a minute without loading them into memory.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    from openpyxl import load_workbook

    source, out_dir = Path(sys.argv[1]).expanduser(), Path(sys.argv[2]).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    for sheet in load_workbook(source, read_only=True, data_only=True).worksheets:
        target = out_dir / f"{sheet.title}.csv"
        rows = 0
        with target.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(["" if cell is None else str(cell) for cell in row])
                rows += 1
        print(f"{target}: {rows} rows")


if __name__ == "__main__":
    main()
